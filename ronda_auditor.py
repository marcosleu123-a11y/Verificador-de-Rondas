import argparse
import base64
import csv
import json
import mimetypes
import os
import re
import statistics
import tempfile
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from PIL import Image, ImageOps, UnidentifiedImageError

from postgres_writer import ExecucaoAuditoria, salvar_auditoria_postgres


GRUPO_VERMELHO = "vermelho"
GRUPO_AMARELO = "amarelo"
GRUPO_VERDE = "verde"
GRUPO_SEM_COMPROVACAO = "sem_comprovacao"
SQL_COLUNA_SEGURA_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*(\.[A-Za-z_][A-Za-z0-9_]*)?$")


def carregar_env(caminho: Path) -> None:
    if not caminho.exists():
        return

    for linha in caminho.read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue

        chave, valor = linha.split("=", 1)
        chave = chave.strip()
        valor = valor.strip().strip('"').strip("'")
        if chave and chave not in os.environ:
            os.environ[chave] = valor


def carregar_env_automatico() -> None:
    env_do_script = Path(__file__).with_name(".env")
    env_da_pasta_atual = Path.cwd() / ".env"

    carregar_env(env_do_script)
    if env_da_pasta_atual.resolve() != env_do_script.resolve():
        carregar_env(env_da_pasta_atual)


@dataclass
class AnaliseImagem:
    grupo: str
    confianca: float
    motivo: str
    brilho_medio: Optional[float] = None
    variacao_visual: Optional[float] = None
    pixels_escuros: Optional[float] = None
    pixels_quase_pretos: Optional[float] = None
    nitidez_aproximada: Optional[float] = None
    largura: Optional[int] = None
    altura: Optional[int] = None


@dataclass
class AnaliseIA:
    grupo: str
    confianca: float
    motivo: str


def resolver_imagem(linha: Dict[str, str], pasta_base: Path) -> Tuple[Optional[Path], Optional[str]]:
    image_path = (linha.get("image_path") or "").strip()
    image_url = (linha.get("image_url") or "").strip()

    if image_path:
        path = Path(image_path)
        if not path.is_absolute():
            path = pasta_base / path
        return path, None

    if image_url:
        try:
            suffix = Path(image_url.split("?")[0]).suffix or ".jpg"
            destino = Path(tempfile.gettempdir()) / f"ronda_auditor_{abs(hash(image_url))}{suffix}"
            if not destino.exists():
                urllib.request.urlretrieve(image_url, destino)
            return destino, None
        except Exception as exc:
            return None, f"erro ao baixar imagem: {exc}"

    return None, "sem comprovacao: image_path e image_url vazios"


def carregar_imagem(caminho: Path) -> Image.Image:
    with Image.open(caminho) as img:
        return img.convert("RGB").copy()


def calcular_nitidez_aproximada(pixels: Iterable[int], largura: int, altura: int) -> float:
    valores = list(pixels)
    if largura < 2 or altura < 2:
        return 0.0

    diferencas = []
    for y in range(altura - 1):
        inicio = y * largura
        proxima_linha = (y + 1) * largura
        for x in range(largura - 1):
            atual = valores[inicio + x]
            direita = valores[inicio + x + 1]
            abaixo = valores[proxima_linha + x]
            diferencas.append(abs(atual - direita))
            diferencas.append(abs(atual - abaixo))

    return float(statistics.mean(diferencas)) if diferencas else 0.0


def analisar_imagem_local(caminho: Optional[Path], erro_previo: Optional[str] = None) -> AnaliseImagem:
    if erro_previo:
        if "sem comprovacao" in erro_previo:
            return AnaliseImagem(GRUPO_SEM_COMPROVACAO, 1.0, erro_previo)
        return AnaliseImagem(GRUPO_VERMELHO, 1.0, erro_previo)

    if caminho is None:
        return AnaliseImagem(GRUPO_SEM_COMPROVACAO, 1.0, "imagem nao informada")

    if not caminho.exists():
        return AnaliseImagem(GRUPO_VERMELHO, 1.0, f"arquivo nao encontrado: {caminho}")

    try:
        img = carregar_imagem(caminho)
    except (UnidentifiedImageError, OSError) as exc:
        return AnaliseImagem(GRUPO_VERMELHO, 1.0, f"arquivo invalido ou corrompido: {exc}")

    largura, altura = img.size
    if largura < 80 or altura < 80:
        return AnaliseImagem(
            GRUPO_AMARELO,
            0.82,
            "imagem com resolucao muito baixa",
            largura=largura,
            altura=altura,
        )

    cinza = ImageOps.grayscale(img)
    cinza.thumbnail((256, 256))
    if hasattr(cinza, "get_flattened_data"):
        pixels = list(cinza.get_flattened_data())
    else:
        pixels = list(cinza.getdata())

    brilho_medio = float(statistics.mean(pixels))
    variacao_visual = float(statistics.pstdev(pixels))
    pixels_escuros = sum(1 for pixel in pixels if pixel < 35) / len(pixels)
    pixels_quase_pretos = sum(1 for pixel in pixels if pixel < 12) / len(pixels)
    nitidez_aproximada = calcular_nitidez_aproximada(pixels, cinza.width, cinza.height)

    metricas = {
        "brilho_medio": brilho_medio,
        "variacao_visual": variacao_visual,
        "pixels_escuros": pixels_escuros,
        "pixels_quase_pretos": pixels_quase_pretos,
        "nitidez_aproximada": nitidez_aproximada,
        "largura": largura,
        "altura": altura,
    }

    if brilho_medio < 18 and variacao_visual < 8 and pixels_escuros > 0.90:
        return AnaliseImagem(
            GRUPO_VERMELHO,
            0.99,
            "imagem praticamente preta, com brilho e variacao visual muito baixos",
            **metricas,
        )

    if pixels_quase_pretos > 0.95 and variacao_visual < 12:
        return AnaliseImagem(
            GRUPO_VERMELHO,
            0.98,
            "quase todos os pixels estao pretos",
            **metricas,
        )

    if brilho_medio < 30 and pixels_escuros > 0.80:
        return AnaliseImagem(
            GRUPO_VERMELHO,
            0.93,
            "imagem muito escura, sem evidencia visual aproveitavel",
            **metricas,
        )

    if brilho_medio < 45 or pixels_escuros > 0.65:
        return AnaliseImagem(
            GRUPO_AMARELO,
            0.78,
            "imagem escura; pode ser evidencia invalida ou ambiente com pouca luz",
            **metricas,
        )

    if variacao_visual < 10:
        return AnaliseImagem(
            GRUPO_AMARELO,
            0.74,
            "imagem quase uniforme, com pouca informacao visual",
            **metricas,
        )

    if nitidez_aproximada < 2.5 and variacao_visual < 25:
        return AnaliseImagem(
            GRUPO_AMARELO,
            0.70,
            "imagem possivelmente borrada ou sem detalhes suficientes",
            **metricas,
        )

    return AnaliseImagem(
        GRUPO_VERDE,
        0.88,
        "imagem passou nos criterios tecnicos basicos",
        **metricas,
    )


def deve_chamar_ia(grupo_local: str, modo: str) -> bool:
    if modo == "todos":
        return True
    if grupo_local == GRUPO_SEM_COMPROVACAO:
        return False
    if modo == "amarelo":
        return grupo_local == GRUPO_AMARELO
    if modo == "amarelo-verde":
        return grupo_local in {GRUPO_AMARELO, GRUPO_VERDE}
    return False


def encode_image_data_url(caminho: Path) -> str:
    mime_type = mimetypes.guess_type(caminho.name)[0] or "image/jpeg"
    conteudo = base64.b64encode(caminho.read_bytes()).decode("utf-8")
    return f"data:{mime_type};base64,{conteudo}"


def analisar_com_ia(caminho: Path, justificativa: str, analise_local: AnaliseImagem) -> Optional[AnaliseIA]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        return None

    from openai import OpenAI

    client = OpenAI(api_key=api_key)
    modelo = os.getenv("OPENAI_MODEL", "gpt-4.1")
    prompt = f"""
Voce e um auditor de evidencias de rondas operacionais.
Classifique a imagem em um dos grupos: vermelho, amarelo ou verde.

Use vermelho quando a imagem for preta, escura demais, vazia, incoerente como evidencia, ou claramente inutil.
Use amarelo quando houver duvida relevante.
Use verde somente quando a imagem parecer uma evidencia visual aproveitavel.

Justificativa informada pelo colaborador: {justificativa or "nao informada"}

Analise tecnica local:
- grupo: {analise_local.grupo}
- motivo: {analise_local.motivo}
- brilho medio: {analise_local.brilho_medio}
- variacao visual: {analise_local.variacao_visual}
- pixels escuros: {analise_local.pixels_escuros}

Responda apenas em JSON valido neste formato:
{{"grupo":"vermelho|amarelo|verde","confianca":0.0,"motivo":"explicacao curta"}}
""".strip()

    response = client.responses.create(
        model=modelo,
        input=[
            {
                "role": "user",
                "content": [
                    {"type": "input_text", "text": prompt},
                    {"type": "input_image", "image_url": encode_image_data_url(caminho), "detail": "low"},
                ],
            }
        ],
    )

    dados = extrair_json(response.output_text)
    grupo = str(dados.get("grupo", "")).lower().strip()
    if grupo not in {GRUPO_VERMELHO, GRUPO_AMARELO, GRUPO_VERDE}:
        raise ValueError(f"grupo de IA invalido: {grupo}")

    return AnaliseIA(
        grupo=grupo,
        confianca=float(dados.get("confianca", 0.0)),
        motivo=str(dados.get("motivo", "")).strip() or "IA nao informou motivo",
    )


def extrair_json(texto: str) -> Dict[str, object]:
    texto = texto.strip()
    try:
        return json.loads(texto)
    except json.JSONDecodeError:
        inicio = texto.find("{")
        fim = texto.rfind("}")
        if inicio == -1 or fim == -1 or fim <= inicio:
            raise
        return json.loads(texto[inicio : fim + 1])


def combinar_resultados(local: AnaliseImagem, ia: Optional[AnaliseIA]) -> Tuple[str, float, str]:
    if ia is None:
        return local.grupo, local.confianca, local.motivo

    if local.grupo == GRUPO_SEM_COMPROVACAO:
        return local.grupo, local.confianca, local.motivo

    if local.grupo == GRUPO_VERMELHO and local.confianca >= 0.93:
        return local.grupo, local.confianca, f"{local.motivo}; IA registrada como apoio: {ia.grupo} - {ia.motivo}"

    if ia.grupo == local.grupo:
        confianca = max(local.confianca, ia.confianca)
        return ia.grupo, confianca, f"{local.motivo}; IA confirmou: {ia.motivo}"

    if GRUPO_VERMELHO in {local.grupo, ia.grupo}:
        return GRUPO_AMARELO, 0.86, f"divergencia entre analise tecnica ({local.grupo}) e IA ({ia.grupo}); revisar"

    return ia.grupo, ia.confianca, f"IA ajustou classificacao: {ia.motivo}"


def salvar_resultado(linhas_saida: List[Dict[str, object]], saida: Path) -> None:
    saida.parent.mkdir(parents=True, exist_ok=True)
    if saida.suffix.lower() == ".xlsx":
        salvar_xlsx(linhas_saida, saida)
        return

    campos = list(linhas_saida[0].keys()) if linhas_saida else []
    with saida.open("w", encoding="utf-8", newline="") as arquivo_saida:
        escritor = csv.DictWriter(arquivo_saida, fieldnames=campos)
        escritor.writeheader()
        escritor.writerows(linhas_saida)


def salvar_xlsx(linhas_saida: List[Dict[str, object]], saida: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = saida.stem[:31] or "Auditoria"

    colunas = list(linhas_saida[0].keys()) if linhas_saida else []
    sheet.append(colunas)
    for linha in linhas_saida:
        sheet.append([linha.get(coluna, "") for coluna in colunas])

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    cores = {
        GRUPO_VERDE: "DCFCE7",
        GRUPO_AMARELO: "FEF3C7",
        GRUPO_VERMELHO: "FEE2E2",
        GRUPO_SEM_COMPROVACAO: "E5E7EB",
    }
    if "grupo" in colunas:
        indice_grupo = colunas.index("grupo") + 1
        for row in range(2, sheet.max_row + 1):
            grupo = str(sheet.cell(row=row, column=indice_grupo).value or "").lower()
            fill = PatternFill("solid", fgColor=cores.get(grupo, "FFFFFF"))
            for col in range(1, sheet.max_column + 1):
                sheet.cell(row=row, column=col).fill = fill

    sheet.freeze_panes = "A2"
    if colunas:
        sheet.auto_filter.ref = sheet.dimensions

    for coluna in range(1, sheet.max_column + 1):
        letra = get_column_letter(coluna)
        maior = 12
        for cell in sheet[letra]:
            maior = max(maior, len(str(cell.value or "")) + 2)
        sheet.column_dimensions[letra].width = min(maior, 55)

    workbook.save(saida)


def auditar_linhas(linhas: List[Dict[str, object]], pasta_base: Path, saida: Path, usar_ia: bool, ia_em: str) -> List[Dict[str, object]]:
    linhas_saida = []
    for linha_original in linhas:
        linha = {chave: "" if valor is None else str(valor) for chave, valor in linha_original.items()}
        caminho, erro = resolver_imagem(linha, pasta_base)
        analise_local = analisar_imagem_local(caminho, erro)
        analise_ia = None
        erro_ia = ""

        if usar_ia and caminho and caminho.exists() and deve_chamar_ia(analise_local.grupo, ia_em):
            if not os.getenv("OPENAI_API_KEY"):
                erro_ia = "OPENAI_API_KEY nao configurada"
            else:
                try:
                    analise_ia = analisar_com_ia(caminho, linha.get("justificativa", ""), analise_local)
                except Exception as exc:
                    erro_ia = str(exc)

        grupo_final, confianca_final, motivo_final = combinar_resultados(analise_local, analise_ia)

        linhas_saida.append(
            {
                **linha,
                "grupo": grupo_final,
                "confianca": f"{confianca_final:.2f}",
                "motivo": motivo_final,
                "grupo_local": analise_local.grupo,
                "confianca_local": f"{analise_local.confianca:.2f}",
                "motivo_local": analise_local.motivo,
                "brilho_medio": formatar_numero(analise_local.brilho_medio),
                "variacao_visual": formatar_numero(analise_local.variacao_visual),
                "pixels_escuros": formatar_numero(analise_local.pixels_escuros),
                "pixels_quase_pretos": formatar_numero(analise_local.pixels_quase_pretos),
                "nitidez_aproximada": formatar_numero(analise_local.nitidez_aproximada),
                "largura": analise_local.largura or "",
                "altura": analise_local.altura or "",
                "grupo_ia": analise_ia.grupo if analise_ia else "",
                "confianca_ia": f"{analise_ia.confianca:.2f}" if analise_ia else "",
                "motivo_ia": analise_ia.motivo if analise_ia else "",
                "erro_ia": erro_ia,
            }
        )

    salvar_resultado(linhas_saida, saida)
    return linhas_saida


def auditar_csv(entrada: Path, saida: Path, usar_ia: bool, ia_em: str) -> List[Dict[str, object]]:
    with entrada.open("r", encoding="utf-8-sig", newline="") as arquivo_entrada:
        amostra = arquivo_entrada.read(4096)
        arquivo_entrada.seek(0)
        dialecto = csv.Sniffer().sniff(amostra, delimiters=",;") if amostra else csv.excel
        leitor = csv.DictReader(arquivo_entrada, dialect=dialecto)
        if not leitor.fieldnames:
            raise ValueError("CSV de entrada esta vazio ou sem cabecalho")
        return auditar_linhas(list(leitor), entrada.parent, saida, usar_ia, ia_em)


def montar_intervalo_datas(data: Optional[str], data_inicio: Optional[str], data_fim: Optional[str]) -> Tuple[datetime, datetime]:
    if data:
        inicio = datetime.strptime(data, "%Y-%m-%d")
        return inicio, inicio + timedelta(days=1)

    if not data_inicio or not data_fim:
        raise ValueError("informe --data ou informe --data-inicio e --data-fim")

    inicio = datetime.strptime(data_inicio, "%Y-%m-%d")
    fim_inclusivo = datetime.strptime(data_fim, "%Y-%m-%d")
    return inicio, fim_inclusivo + timedelta(days=1)


def montar_conexao_sql(args: argparse.Namespace) -> str:
    driver = os.getenv("SQL_DRIVER", "ODBC Driver 17 for SQL Server")
    partes = [
        f"DRIVER={{{driver}}}",
        f"SERVER={args.sql_server}",
        f"DATABASE={args.sql_database}",
        "TrustServerCertificate=yes",
        "Encrypt=no",
    ]

    if args.sql_auth == "windows":
        partes.append("Trusted_Connection=yes")
    else:
        usuario = args.sql_user or os.getenv("SQL_USER")
        senha = os.getenv(args.sql_password_env)
        if not usuario or not senha:
            raise ValueError(f"configure SQL_USER e {args.sql_password_env}, ou informe --sql-user")
        partes.extend([f"UID={usuario}", f"PWD={senha}"])

    return ";".join(partes)


def campo_periodo_sql(campo_periodo: str) -> str:
    campos = {
        "disponibilizacao": "i.Execucao_Disponibilizacao",
        "tarefa_disponibilizacao": "t.Disponibilizacao",
        "inicio": "t.Inicio",
        "termino": "t.Termino",
        "prazo": "t.Prazo",
        "inicio_real": "t.InicioReal",
        "execucao": "i.Data_Inicio",
    }
    return campos[campo_periodo]


def coluna_cr_sql(coluna_cr: Optional[str]) -> str:
    if not coluna_cr:
        return "CAST(NULL AS varchar(100))"
    coluna_cr = coluna_cr.strip()
    if not SQL_COLUNA_SEGURA_RE.match(coluna_cr):
        raise ValueError(
            "--sql-coluna-cr deve ser um identificador simples, como CR, t.CR ou t.Contrato_CR"
        )
    return coluna_cr


def buscar_linhas_sql(args: argparse.Namespace) -> List[Dict[str, object]]:
    import pyodbc

    data_inicio, data_fim = montar_intervalo_datas(args.data, args.data_inicio, args.data_fim)
    campo_periodo = campo_periodo_sql(args.campo_periodo)
    campo_cr = coluna_cr_sql(args.sql_coluna_cr)
    order_by = f"{campo_periodo} DESC, i.Numero DESC"
    if campo_periodo != "i.Data_Inicio":
        order_by = f"{campo_periodo} DESC, i.Data_Inicio DESC"

    query = f"""
WITH InicioNao AS (
    SELECT
        Numero,
        Tarefa_id,
        Colaborador,
        MIN(Data) AS Data_Inicio,
        MIN(Disponibilizacao) AS Execucao_Disponibilizacao
    FROM dbo.Execucao
    WHERE Descricao = 'INICIAR A RONDA?'
      AND Conteudo COLLATE Latin1_General_CI_AI IN ('Nao', 'NAO', 'nao')
    GROUP BY
        Numero,
        Tarefa_id,
        Colaborador
),
Justificativa AS (
    SELECT
        Numero,
        Tarefa_id,
        MAX(CASE
            WHEN Conteudo IS NOT NULL
             AND Conteudo NOT LIKE 'http%'
            THEN Conteudo
        END) AS justificativa,
        MAX(CASE
            WHEN Conteudo LIKE 'http%'
            THEN Conteudo
        END) AS image_url,
        MAX(Data) AS data_justificativa
    FROM dbo.Execucao
    WHERE Descricao = 'Justifique e evidencie com foto.'
    GROUP BY
        Numero,
        Tarefa_id
),
EstruturaExecucao AS (
    SELECT
        Numero,
        Tarefa_id,
        MAX(CASE
            WHEN SKESTRUTURA IS NOT NULL
             AND CONVERT(varchar(100), SKESTRUTURA) <> '-1'
            THEN CONVERT(varchar(100), SKESTRUTURA)
        END) AS Execucao_SKESTRUTURA,
        MAX(CASE
            WHEN QRCode IS NOT NULL
             AND LTRIM(RTRIM(QRCode)) <> ''
             AND QRCode COLLATE Latin1_General_CI_AI NOT IN ('Nao', 'NAO', 'nao')
            THEN QRCode
        END) AS Execucao_QRCode
    FROM dbo.Execucao
    GROUP BY
        Numero,
        Tarefa_id
)
SELECT
    i.Numero AS atividade_id,
    i.Tarefa_id AS tarefa_id,
    COALESCE({campo_cr}, e_exec.ID_CR, e_tarefa.ID_CR) AS contrato_cr,
    COALESCE(e_exec.ID_ESTRUTURA, e_tarefa.ID_ESTRUTURA) AS estrutura_id,
    COALESCE(e_exec.NIVEL_03, e_tarefa.NIVEL_03) AS nivel_03,
    COALESCE(e_exec.NIVEL_04, e_tarefa.NIVEL_04) AS nivel_04,
    COALESCE(e_exec.ANDAR, e_tarefa.ANDAR) AS andar,
    COALESCE(e_exec.LOCAL, e_tarefa.LOCAL) AS local,
    COALESCE(e_exec.AMBIENTE, e_tarefa.AMBIENTE) AS ambiente,
    COALESCE(ee.Execucao_QRCode, e_exec.QRCODE, e_tarefa.QRCODE) AS qrcode,
    i.Colaborador AS colaborador,
    {campo_periodo} AS data,
    i.Data_Inicio AS data_execucao_inicio,
    i.Execucao_Disponibilizacao AS execucao_disponibilizacao,
    t.Disponibilizacao AS tarefa_disponibilizacao,
    t.Inicio AS tarefa_inicio,
    t.Termino AS tarefa_termino,
    t.Prazo AS tarefa_prazo,
    t.Nome AS tarefa_nome,
    t.Checklist_Nome AS checklist_nome,
    t.Checklist_Descricao AS checklist_descricao,
    j.justificativa,
    j.image_url,
    j.data_justificativa
FROM InicioNao i
LEFT JOIN dbo.TAREFA t
    ON CONVERT(varchar(100), t.tarefa_id) = CONVERT(varchar(100), i.Tarefa_id)
LEFT JOIN EstruturaExecucao ee
    ON ee.Numero = i.Numero
   AND ee.Tarefa_id = i.Tarefa_id
LEFT JOIN dbo.D_ESTRUTURA e_exec
    ON CONVERT(varchar(100), e_exec.SKESTRUTURA) = CONVERT(varchar(100), ee.Execucao_SKESTRUTURA)
LEFT JOIN dbo.D_ESTRUTURA e_tarefa
    ON CONVERT(varchar(100), e_tarefa.ID_ESTRUTURA) = CONVERT(varchar(100), t.estrutura_id)
LEFT JOIN Justificativa j
    ON j.Numero = i.Numero
   AND j.Tarefa_id = i.Tarefa_id
WHERE {campo_periodo} >= ?
  AND {campo_periodo} < ?
ORDER BY {order_by};
"""

    with pyodbc.connect(montar_conexao_sql(args), timeout=30) as conexao:
        cursor = conexao.cursor()
        cursor.execute(query, data_inicio, data_fim)
        colunas = [coluna[0] for coluna in cursor.description]
        return [dict(zip(colunas, linha)) for linha in cursor.fetchall()]


def formatar_numero(valor: Optional[float]) -> str:
    return "" if valor is None else f"{valor:.4f}"


def periodo_execucao(args: argparse.Namespace) -> Tuple[Optional[datetime], Optional[datetime]]:
    if args.data or (args.data_inicio and args.data_fim):
        inicio, fim_exclusivo = montar_intervalo_datas(args.data, args.data_inicio, args.data_fim)
        return inicio, fim_exclusivo - timedelta(days=1)
    return None, None


def criar_parser() -> argparse.ArgumentParser:
    carregar_env_automatico()

    parser = argparse.ArgumentParser(description="Audita evidencias de rondas justificadas.")
    parser.add_argument("--saida", required=True, help="CSV de resultado da auditoria")
    parser.add_argument("--entrada", help="CSV com atividades justificadas")
    parser.add_argument("--salvar-postgres", action="store_true", help="Grava rondas e analise do script no Postgres")
    parser.add_argument(
        "--substituir-periodo",
        action="store_true",
        help="Antes de salvar no Postgres, remove execucoes anteriores do mesmo periodo/origem/campo",
    )
    parser.add_argument(
        "--substituir-sobreposicoes",
        action="store_true",
        help="Antes de salvar no Postgres, remove execucoes antigas com periodo sobreposto ao novo",
    )
    parser.add_argument("--sql-server", default=os.getenv("SQL_SERVER"), help="Servidor SQL Server. Exemplo: seu_servidor,1433")
    parser.add_argument("--sql-database", default=os.getenv("SQL_DATABASE"), help="Nome do banco SQL Server. Exemplo: nome_do_banco")
    parser.add_argument(
        "--sql-auth",
        choices=["sql", "windows"],
        default=os.getenv("SQL_AUTH", "sql").lower(),
        help="Tipo de autenticacao no SQL Server",
    )
    parser.add_argument("--sql-user", default=os.getenv("SQL_USER"), help="Usuario SQL. Tambem pode usar a variavel SQL_USER")
    parser.add_argument("--sql-password-env", default="SQL_PASSWORD", help="Nome da variavel de ambiente com a senha SQL")
    parser.add_argument(
        "--sql-coluna-cr",
        default=os.getenv("SQL_COLUNA_CR"),
        help="Coluna do SQL Server usada como CR do contrato. Exemplo: t.CR ou t.Contrato_CR",
    )
    parser.add_argument("--data", help="Data unica da auditoria, no formato AAAA-MM-DD")
    parser.add_argument("--data-inicio", help="Data inicial, no formato AAAA-MM-DD")
    parser.add_argument("--data-fim", help="Data final inclusiva, no formato AAAA-MM-DD")
    parser.add_argument(
        "--campo-periodo",
        choices=["disponibilizacao", "tarefa_disponibilizacao", "inicio", "termino", "prazo", "inicio_real", "execucao"],
        default=os.getenv("SQL_CAMPO_PERIODO", "disponibilizacao").lower(),
        help="Campo usado para filtrar o periodo no modo banco",
    )
    parser.add_argument("--usar-ia", action="store_true", help="Usa IA de visao quando houver OPENAI_API_KEY")
    parser.add_argument(
        "--ia-em",
        choices=["amarelo", "amarelo-verde", "todos"],
        default="amarelo",
        help="Define quais grupos locais serao enviados para IA",
    )
    return parser


def main() -> None:
    args = criar_parser().parse_args()
    saida = Path(args.saida)

    if args.substituir_periodo and not args.salvar_postgres:
        raise ValueError("--substituir-periodo so pode ser usado junto com --salvar-postgres")
    if args.substituir_sobreposicoes and not args.salvar_postgres:
        raise ValueError("--substituir-sobreposicoes so pode ser usado junto com --salvar-postgres")
    if (args.substituir_periodo or args.substituir_sobreposicoes) and not (
        args.data or (args.data_inicio and args.data_fim)
    ):
        raise ValueError("--substituir-periodo/--substituir-sobreposicoes exige periodo com --data ou --data-inicio/--data-fim")

    if args.entrada:
        entrada = Path(args.entrada)
        linhas_saida = auditar_csv(entrada, saida, args.usar_ia, args.ia_em)
        origem = "csv"
    else:
        if not args.sql_server or not args.sql_database:
            raise ValueError("informe --entrada ou informe --sql-server, --sql-database e a data")
        linhas = buscar_linhas_sql(args)
        print(f"Campo de periodo usado: {args.campo_periodo}")
        print(f"Registros encontrados no banco: {len(linhas)}")
        linhas_saida = auditar_linhas(linhas, Path.cwd(), saida, args.usar_ia, args.ia_em)
        origem = "sql_server"

    if args.salvar_postgres:
        data_inicio, data_fim = periodo_execucao(args)
        execucao_id = salvar_auditoria_postgres(
            linhas_saida,
            ExecucaoAuditoria(
                origem=origem,
                data_inicio=data_inicio.date() if data_inicio else None,
                data_fim=data_fim.date() if data_fim else None,
                campo_periodo=args.campo_periodo if origem == "sql_server" else None,
                usar_ia=args.usar_ia,
                arquivo_excel=saida,
            ),
            substituir_periodo=args.substituir_periodo,
            substituir_sobreposicoes=args.substituir_sobreposicoes,
        )
        print(f"Auditoria salva no Postgres com execucao_id={execucao_id}")

    print(f"Auditoria concluida: {saida}")


if __name__ == "__main__":
    main()
